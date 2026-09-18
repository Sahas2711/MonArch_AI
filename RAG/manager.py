from typing import Optional
from langchain_community.vectorstores import FAISS
from langchain_text_splitters import RecursiveCharacterTextSplitter
from RAG.embeddings import get_embeddings_model
from RAG.retriever import build_bm25_retriever, reciprocal_rank_fusion
from utils.config import EMBEDDING_MODEL, UNSTRUCTURED_API_KEY
from utils.logger import log


class RAGAgentManager:
    """Owns the vector store + keyword index. Singleton instance reused across graph invocations."""

    def __init__(self, embedding_model: str = EMBEDDING_MODEL):
        self.embeddings = get_embeddings_model(embedding_model)
        self.vector_db = FAISS.from_texts(
            ["Monarch RAG baseline — no documents ingested yet."], self.embeddings
        )
        self.all_documents: list = []
        self._bm25_retriever: Optional[object] = None
        self._bm25_dirty = True

    def ingest(self, file_path: str, user_id: Optional[str] = None) -> dict:
        """Load, chunk, tag with metadata, and index text, PDF, DOCX, or image documents."""
        import os
        from langchain_core.documents import Document

        ext = os.path.splitext(file_path)[1].lower()
        docs = []

        # 1. Primary loader: Try UnstructuredLoader
        try:
            from langchain_unstructured import UnstructuredLoader
            loader = UnstructuredLoader(
                file_path=file_path,
                api_key=UNSTRUCTURED_API_KEY,
                partition_via_api=bool(UNSTRUCTURED_API_KEY),
            )
            docs = loader.load()
        except Exception as exc:
            log.info("UnstructuredLoader skipped/failed for %s (%s). Falling back to standard format loader.", file_path, exc)

        # 2. Native Fallback Loaders by extension if UnstructuredLoader produced nothing
        if not docs:
            if ext == ".pdf":
                try:
                    from langchain_community.document_loaders import PyPDFLoader
                    docs = PyPDFLoader(file_path).load()
                except Exception:
                    import pypdf
                    reader = pypdf.PdfReader(file_path)
                    text = "\n\n".join([page.extract_text() or "" for page in reader.pages])
                    docs = [Document(page_content=text, metadata={"source": file_path})]

            elif ext in (".docx", ".doc"):
                try:
                    from langchain_community.document_loaders import Docx2txtLoader
                    docs = Docx2txtLoader(file_path).load()
                except Exception:
                    import docx
                    doc = docx.Document(file_path)
                    text = "\n".join([p.text for p in doc.paragraphs if p.text])
                    docs = [Document(page_content=text, metadata={"source": file_path})]

            elif ext in (".png", ".jpg", ".jpeg"):
                try:
                    import pytesseract
                    from PIL import Image
                    ocr_text = pytesseract.image_to_string(Image.open(file_path))
                    content = f"[Visual Ingest OCR Text]\n{ocr_text}" if ocr_text.strip() else f"[Visual Asset File: {os.path.basename(file_path)}]"
                    docs = [Document(page_content=content, metadata={"source": file_path})]
                except Exception:
                    docs = [Document(page_content=f"[Visual Asset File: {os.path.basename(file_path)}]", metadata={"source": file_path})]

            else:
                try:
                    from langchain_community.document_loaders import TextLoader
                    docs = TextLoader(file_path, encoding="utf-8").load()
                except Exception:
                    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                        docs = [Document(page_content=f.read(), metadata={"source": file_path})]

        splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=120)
        chunks = splitter.split_documents(docs)

        for i, chunk in enumerate(chunks):
            chunk.metadata["user_id"] = user_id
            chunk.metadata["source"] = file_path
            chunk.metadata["chunk_index"] = i

        self.vector_db.add_documents(chunks)
        self.all_documents.extend(chunks)
        self._bm25_dirty = True

        log.info("Ingested %d chunks from %s (user_id=%s)", len(chunks), file_path, user_id)
        return {"status": "success", "chunks_added": len(chunks)}

    def _get_bm25(self):
        if not self.all_documents:
            return None
        if self._bm25_dirty or self._bm25_retriever is None:
            self._bm25_retriever = build_bm25_retriever(self.all_documents)
            self._bm25_dirty = False
        return self._bm25_retriever

    def retrieve(self, query: str, user_id: Optional[str] = None) -> str:
        if not self.all_documents:
            return "(no relevant context found)"

        # Filter active documents by user_id if specified (matching user_id or None)
        valid_docs = self.all_documents
        if user_id:
            user_filtered = [d for d in self.all_documents if d.metadata.get("user_id") in (user_id, None)]
            if user_filtered:
                valid_docs = user_filtered

        # 1. Standard Hybrid Retrieval (FAISS + BM25)
        faiss_retriever = self.vector_db.as_retriever(search_kwargs={"k": 10})
        bm25_retriever = self._get_bm25()

        faiss_docs = faiss_retriever.invoke(query)
        if bm25_retriever is not None:
            bm25_docs = bm25_retriever.invoke(query)
            raw_docs = reciprocal_rank_fusion([bm25_docs, faiss_docs], top_n=5)
        else:
            raw_docs = faiss_docs[:5]

        # Filter retrieved docs to valid_docs / user_id
        if user_id:
            matching_docs = [d for d in raw_docs if d.metadata.get("user_id") in (user_id, None)]
        else:
            matching_docs = raw_docs

        # Exclude the baseline placeholder text
        matching_docs = [d for d in matching_docs if "Monarch RAG baseline" not in d.page_content]

        # 2. Smart Fallback for Broad / Overview / Filename / Conversational Queries
        overview_keywords = {"paper", "document", "file", "pdf", "ingested", "summarize", "about", "overview", "abstract", "research", "title", "that"}
        query_lower = query.lower()
        is_broad_query = any(kw in query_lower for kw in overview_keywords) or len(query.strip().split()) <= 4

        if not matching_docs or is_broad_query:
            # If standard search missed or user asked an overview query, fetch primary chunks from valid_docs
            fallback_chunks = [d for d in valid_docs if "Monarch RAG baseline" not in d.page_content][:5]
            if fallback_chunks:
                seen_texts = set()
                combined = []
                for d in matching_docs + fallback_chunks:
                    if d.page_content not in seen_texts:
                        seen_texts.add(d.page_content)
                        combined.append(d)
                matching_docs = combined[:5]

        if not matching_docs:
            return "(no relevant context found)"

        return "\n\n---\n\n".join(d.page_content for d in matching_docs)


rag_manager = RAGAgentManager()
